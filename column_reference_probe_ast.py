"""
column_reference_probe_ast.py - sqlglot AST 기반 Column Reference Probe

기존 정규식 기반 probe의 한계:
  - false_alias_parse: alias 없이 테이블명 직접 쓴 경우를 unresolved로 오탐 (20건)
  - false_cte: CTE alias를 table_not_found로 오탐 (3건)
  - false_subquery: 서브쿼리 alias 오탐 (1건)

sqlglot 기반 개선:
  - exp.Table에서 alias_map을 정확히 추출 (CTE, 서브쿼리 포함)
  - exp.Column에서 qualifier를 alias_map으로 resolve
  - unqualified column (table='')은 스킵 (false positive 방지)

사용법:
  python column_reference_probe_ast.py

출력:
  results/column_reference_probe_ast_log.json
  results/column_reference_probe_ast_stats.txt
"""

import json
import sqlite3
import re
import os
import openpyxl
from tqdm import tqdm

try:
    import sqlglot
    from sqlglot import exp as sqlglot_exp
except ImportError:
    print("❌ sqlglot이 필요합니다: pip install sqlglot")
    exit(1)

DB_BASE_PATH = "./data/dev_databases"
INPUT_PATH = "./results/failure_corpus_official.json"
OUTPUT_LOG_PATH = "./results/column_reference_probe_ast_log.json"
OUTPUT_STATS_PATH = "./results/column_reference_probe_ast_stats.txt"
FAIL_LABEL_PATH = "./results/error_labeling_233.xlsx"


# ============================================================
# 수작업 라벨 로드
# ============================================================

def load_manual_labels():
    labels = {}
    if not os.path.exists(FAIL_LABEL_PATH):
        return labels
    wb = openpyxl.load_workbook(FAIL_LABEL_PATH)
    ws = wb['오류 라벨링']
    for row in range(2, 235):
        qid = ws.cell(row=row, column=2).value
        label = ws.cell(row=row, column=12).value
        if qid and label:
            labels[qid] = label
    return labels


def extract_main_type(label):
    if not label:
        return 'unknown'
    label = str(label).strip()
    main = label.split('(')[0].strip() if '(' in label else label
    if main == 'COMPLEX':
        match = re.search(r'\(([^)]+)\)', label)
        if match:
            return match.group(1).split(',')[0].strip()
        return 'COMPLEX'
    return main


# ============================================================
# DB 스키마 로드
# ============================================================

def get_all_table_columns(db_id):
    """DB의 모든 table -> columns 매핑"""
    db_path = os.path.join(DB_BASE_PATH, db_id, f"{db_id}.sqlite")
    table_columns = {}
    if not os.path.exists(db_path):
        return table_columns
    try:
        conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
        tables = [row[0] for row in cursor.fetchall()]
        for table in tables:
            cursor.execute(f'PRAGMA table_info([{table}])')
            cols = [row[1].lower() for row in cursor.fetchall()]
            table_columns[table.lower()] = {
                "table_name": table,
                "columns": cols
            }
        conn.close()
    except Exception:
        pass
    return table_columns


def execute_column_probe(db_id, table_name, column_name):
    """SELECT column FROM table LIMIT 1 실행으로 컬럼 존재 검증"""
    db_path = os.path.join(DB_BASE_PATH, db_id, f"{db_id}.sqlite")
    try:
        conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
        cursor = conn.cursor()
        cursor.execute(f'SELECT "{column_name}" FROM "{table_name}" LIMIT 1')
        cursor.fetchall()
        conn.close()
        return True, "ok"
    except Exception as e:
        return False, str(e)


# ============================================================
# sqlglot AST 기반 probe
# ============================================================

def build_alias_map_ast(parsed):
    """sqlglot AST에서 alias → 실제 테이블명 매핑 추출
    
    - CTE alias: WITH temp AS (...) → temp는 가상 테이블로 표시
    - 서브쿼리 alias: (SELECT ...) AS t → t는 가상 테이블
    - 일반 alias: FROM Examination e → e → examination
    - alias 없음: FROM races → races → races
    """
    alias_map = {}
    cte_names = set()

    # CTE 이름 수집 (가상 테이블이라 column 검증 스킵)
    for cte in parsed.find_all(sqlglot_exp.CTE):
        if cte.alias:
            cte_names.add(cte.alias.lower())

    # 서브쿼리 alias 수집
    subquery_aliases = set()
    for subq in parsed.find_all(sqlglot_exp.Subquery):
        if subq.alias:
            subquery_aliases.add(subq.alias.lower())

    # 일반 테이블 alias 추출
    for table in parsed.find_all(sqlglot_exp.Table):
        tname = table.name.lower() if table.name else ''
        alias = table.alias.lower() if table.alias else tname

        if not tname:
            continue

        # CTE/서브쿼리 alias는 가상 테이블로 표시
        if tname in cte_names or alias in cte_names:
            alias_map[alias] = f"__cte__{tname}"
            continue

        alias_map[alias] = tname

    # 서브쿼리 alias도 가상으로 표시
    for sq_alias in subquery_aliases:
        if sq_alias not in alias_map:
            alias_map[sq_alias] = f"__subquery__{sq_alias}"

    return alias_map, cte_names, subquery_aliases


def probe_column_references_ast(db_id, pred_sql):
    """sqlglot AST 기반 column reference 검증
    
    기존 정규식 대비 개선:
    - alias 없이 테이블명 직접 쓴 경우 정확히 처리
    - CTE/서브쿼리 alias 오탐 방지
    - unqualified column 스킵
    
    Returns:
        dict with invalid_refs, unresolved_refs, has_invalid_ref, has_unresolved_ref
    """
    table_columns = get_all_table_columns(db_id)

    try:
        parsed = sqlglot.parse_one(pred_sql, dialect="sqlite", error_level=sqlglot.ErrorLevel.IGNORE)
    except Exception as e:
        return {
            "parse_error": str(e),
            "invalid_refs": [],
            "unresolved_refs": [],
            "has_invalid_ref": False,
            "has_unresolved_ref": False,
        }

    alias_map, cte_names, subquery_aliases = build_alias_map_ast(parsed)

    invalid_refs = []
    unresolved_refs = []
    seen = set()

    for col in parsed.find_all(sqlglot_exp.Column):
        qualifier = col.table.lower() if col.table else ''
        column = col.name.lower() if col.name else ''

        if not qualifier or not column:
            continue  # unqualified column 스킵

        key = (qualifier, column)
        if key in seen:
            continue
        seen.add(key)

        if qualifier not in alias_map:
            unresolved_refs.append({
                "qualifier": qualifier,
                "column": column,
                "issue": "unresolved_qualifier"
            })
            continue

        resolved_table = alias_map[qualifier]

        # CTE/서브쿼리 alias는 검증 스킵
        if resolved_table.startswith("__cte__") or resolved_table.startswith("__subquery__"):
            continue

        if resolved_table not in table_columns:
            invalid_refs.append({
                "qualifier": qualifier,
                "resolved_table": resolved_table,
                "column": column,
                "issue": "table_not_in_schema",
            })
            continue

        actual_cols = table_columns[resolved_table]["columns"]
        actual_table_name = table_columns[resolved_table]["table_name"]

        if column not in actual_cols:
            # 다른 테이블에 있는지 확인
            candidate_tables = [
                info["table_name"]
                for t_l, info in table_columns.items()
                if column in info["columns"] and t_l != resolved_table
            ]

            # execute-based 검증
            ok, msg = execute_column_probe(db_id, actual_table_name, column)

            invalid_refs.append({
                "qualifier": qualifier,
                "resolved_table": actual_table_name,
                "column": column,
                "issue": "column_not_in_resolved_table",
                "candidate_tables": candidate_tables,
                "exec_probe": ok,
                "exec_message": msg,
            })
        else:
            # PRAGMA 기준 정상 — execute 확인
            ok, msg = execute_column_probe(db_id, actual_table_name, column)
            if not ok:
                invalid_refs.append({
                    "qualifier": qualifier,
                    "resolved_table": actual_table_name,
                    "column": column,
                    "issue": "execute_failed",
                    "exec_message": msg,
                })

    return {
        "alias_map": alias_map,
        "cte_names": list(cte_names),
        "subquery_aliases": list(subquery_aliases),
        "invalid_refs": invalid_refs,
        "unresolved_refs": unresolved_refs,
        "has_invalid_ref": len(invalid_refs) > 0,
        "has_unresolved_ref": len(unresolved_refs) > 0,
    }


# ============================================================
# 메인
# ============================================================

def main():
    print("🔍 sqlglot AST 기반 Column Reference Probe 시작")

    with open(INPUT_PATH, 'r', encoding='utf-8') as f:
        samples = json.load(f)

    manual_labels = load_manual_labels()
    print(f"📂 실패 코퍼스: {len(samples)}건")
    print(f"📂 수작업 라벨: {len(manual_labels)}건\n")

    logs = []

    for sample in tqdm(samples):
        db_id = sample['db_id']
        pred_sql = sample['pred_sql']
        qid = sample.get('question_id')

        probe = probe_column_references_ast(db_id, pred_sql)
        manual_label = manual_labels.get(qid)

        logs.append({
            "question_id": qid,
            "db_id": db_id,
            "difficulty": sample.get('difficulty'),
            "error_type": sample.get('error_type'),
            "manual_label": manual_label,
            "pred_sql": pred_sql,
            "column_probe": probe,
        })

    with open(OUTPUT_LOG_PATH, 'w', encoding='utf-8') as f:
        json.dump(logs, f, ensure_ascii=False, indent=2)

    # ============================================================
    # 통계
    # ============================================================

    total = len(logs)
    invalid_count = sum(1 for x in logs if x['column_probe'].get('has_invalid_ref'))
    unresolved_count = sum(1 for x in logs if x['column_probe'].get('has_unresolved_ref'))
    parse_error_count = sum(1 for x in logs if x['column_probe'].get('parse_error'))

    # 라벨별 분포
    label_stats = {}
    for x in logs:
        label = extract_main_type(x.get('manual_label'))
        if label not in label_stats:
            label_stats[label] = {'total': 0, 'invalid': 0, 'unresolved': 0}
        label_stats[label]['total'] += 1
        if x['column_probe'].get('has_invalid_ref'):
            label_stats[label]['invalid'] += 1
        if x['column_probe'].get('has_unresolved_ref'):
            label_stats[label]['unresolved'] += 1

    # issue 유형별
    issue_counts = {}
    candidate_count = 0
    for x in logs:
        for ref in x['column_probe'].get('invalid_refs', []):
            issue = ref.get('issue', 'unknown')
            issue_counts[issue] = issue_counts.get(issue, 0) + 1
            if ref.get('candidate_tables'):
                candidate_count += 1

    stats_lines = []
    stats_lines.append("=" * 65)
    stats_lines.append("📊 sqlglot AST Column Reference Probe 통계")
    stats_lines.append("=" * 65)
    stats_lines.append(f"전체: {total}건")
    stats_lines.append(f"파싱 오류: {parse_error_count}건")
    stats_lines.append(f"invalid ref 있음: {invalid_count}건 ({invalid_count/total*100:.1f}%)")
    stats_lines.append(f"unresolved qualifier 있음: {unresolved_count}건 ({unresolved_count/total*100:.1f}%)")
    stats_lines.append(f"candidate_tables 있음: {candidate_count}건")

    stats_lines.append(f"\n수작업 라벨별 invalid_ref 발생률:")
    stats_lines.append(f"  {'라벨':<20} {'전체':>6} {'invalid':>8} {'비율':>8} {'unresolved':>12}")
    stats_lines.append("  " + "-" * 58)
    for label, s in sorted(label_stats.items(), key=lambda x: -x[1]['total']):
        rate = s['invalid'] / s['total'] * 100 if s['total'] > 0 else 0
        stats_lines.append(
            f"  {label:<20} {s['total']:>6} {s['invalid']:>8} {rate:>7.1f}% {s['unresolved']:>12}"
        )

    stats_lines.append(f"\ninvalid ref issue 유형:")
    for issue, cnt in sorted(issue_counts.items(), key=lambda x: -x[1]):
        stats_lines.append(f"  {issue}: {cnt}건")

    # 기존 정규식 결과와 비교
    stats_lines.append(f"\n=== 기존 정규식 대비 비교 ===")
    stats_lines.append(f"  기존 invalid: 19건 (7.3%) → AST: {invalid_count}건 ({invalid_count/total*100:.1f}%)")
    stats_lines.append(f"  기존 unresolved: 23건 (8.8%) → AST: {unresolved_count}건 ({unresolved_count/total*100:.1f}%)")
    stats_lines.append(f"  기존 false_alias_parse 20건 → AST에서 제거됨 (예상)")

    stats_lines.append("=" * 65)

    stats_text = "\n".join(stats_lines)
    print(stats_text)

    with open(OUTPUT_STATS_PATH, 'w', encoding='utf-8') as f:
        f.write(stats_text)

    print(f"\n✅ 로그: {OUTPUT_LOG_PATH}")
    print(f"✅ 통계: {OUTPUT_STATS_PATH}")


if __name__ == "__main__":
    main()